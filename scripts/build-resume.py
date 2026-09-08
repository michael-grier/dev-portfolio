"""Build the editable resume from the content used by the Next.js resume page.

Requires python-docx. Run from any directory; outputs public/Michael_Grier_Resume.docx.
The bundled fonts are embedded so Word can preserve the portfolio typography.
"""

import argparse
import json
from pathlib import Path
from uuid import uuid4

from lxml import etree
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.opc.packuri import PackURI
from docx.opc.part import Part
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
FONT = "Familjen Grotesk"
INK = "1C1C1C"
BLUE = "0078BF"
MUTED = "5F5C57"


def set_font(style, size, color=INK, bold=False):
    """Set explicit font properties instead of inheriting Word's theme fonts."""
    style.font.name = FONT
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor.from_string(color)
    style.font.bold = bold
    style.font.italic = False
    style.font.underline = False
    for border in style.element.findall(".//" + qn("w:pBdr")):
        border.getparent().remove(border)
    fonts = style.element.get_or_add_rPr().get_or_add_rFonts()
    for key in list(fonts.attrib):
        if "Theme" in key:
            del fonts.attrib[key]
    for script in ("ascii", "hAnsi", "eastAsia", "cs"):
        fonts.set(qn(f"w:{script}"), FONT)


def rule(paragraph, color, size=6):
    """Draw a light section rule that is independent of the text reading order."""
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    for name, value in {"val": "single", "sz": str(size), "space": "5", "color": color}.items():
        bottom.set(qn(f"w:{name}"), value)
    borders.append(bottom)
    paragraph._p.get_or_add_pPr().append(borders)


def link(paragraph, label, url):
    """Add a clickable text link while preserving the paragraph's typography."""
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True))
    run = OxmlElement("w:r")
    props = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), BLUE)
    props.append(color)
    run.append(props)
    text = OxmlElement("w:t")
    text.text = label
    run.append(text)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def embed_fonts(document):
    """Embed full licensed fonts using the OOXML font obfuscation format."""
    part = document.part.part_related_by(RT.FONT_TABLE)
    table = parse_xml(part.blob)
    entry = OxmlElement("w:font")
    entry.set(qn("w:name"), FONT)
    for variant in ("Regular", "Bold"):
        key = uuid4()
        data = bytearray((ROOT / "scripts" / "resume-assets" / f"FamiljenGrotesk-{variant}.ttf").read_bytes())
        # OOXML XORs the first 32 bytes with the reversed GUID bytes twice.
        mask = key.bytes[::-1]
        for index in range(32):
            data[index] ^= mask[index % 16]
        font_part = Part(
            PackURI(f"/word/fonts/FamiljenGrotesk-{variant}.odttf"),
            "application/vnd.openxmlformats-officedocument.obfuscatedFont",
            bytes(data),
            document.part.package,
        )
        embedded = OxmlElement(f"w:embed{variant}")
        embedded.set(qn("r:id"), part.relate_to(font_part, RT.FONT))
        embedded.set(qn("w:fontKey"), "{" + str(key).upper() + "}")
        embedded.set(qn("w:subsetted"), "false")
        entry.append(embedded)
    table.append(entry)
    part._blob = etree.tostring(table)
    settings = OxmlElement("w:embedTrueTypeFonts")
    document.settings.element.append(settings)


def heading(document, text):
    """Start a section and keep its label attached to the following content."""
    paragraph = document.add_paragraph(text, "Heading 1")
    rule(paragraph, BLUE)
    return paragraph


def lightburst(paragraph, width):
    """Anchor corner artwork behind text so it cannot change pagination or reading order."""
    picture = paragraph.add_run().add_picture(
        str(ROOT / "scripts/resume-assets/lightburst.png"), width=Inches(width)
    )
    inline = picture._inline
    anchor = OxmlElement("wp:anchor")
    for key, value in {"distT": "0", "distB": "0", "distL": "0", "distR": "0",
                       "simplePos": "0", "relativeHeight": "0", "behindDoc": "1",
                       "locked": "0", "layoutInCell": "1", "allowOverlap": "1"}.items():
        anchor.set(key, value)
    simple = OxmlElement("wp:simplePos")
    simple.set("x", "0")
    simple.set("y", "0")
    anchor.append(simple)
    for axis, offset in [("H", Inches(8.5 - width)), ("V", 0)]:
        position = OxmlElement(f"wp:position{axis}")
        position.set("relativeFrom", "page")
        value = OxmlElement("wp:posOffset")
        value.text = str(offset)
        position.append(value)
        anchor.append(position)
    anchor.append(inline.find(qn("wp:extent")))
    anchor.append(OxmlElement("wp:wrapNone"))
    for name in ("wp:docPr", "wp:cNvGraphicFramePr", "a:graphic"):
        anchor.append(inline.find(qn(name)))
    inline.getparent().replace(inline, anchor)


def build(output):
    """Lay out a two-page resume with experience first and projects on page two."""
    data = json.loads((ROOT / "src/content/resume.json").read_text())
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.48)
    section.bottom_margin = Inches(0.48)
    section.left_margin = section.right_margin = Inches(0.62)
    section.header_distance = section.footer_distance = Inches(0.22)
    section.different_first_page_header_footer = True
    normal = doc.styles["Normal"]
    set_font(normal, 10.5)
    normal.paragraph_format.line_spacing = 1.04
    normal.paragraph_format.space_after = Pt(4)
    normal.paragraph_format.widow_control = True
    for name, size, color, bold in [
        ("Title", 34, INK, True),
        ("Subtitle", 15, MUTED, False),
        ("Heading 1", 13, BLUE, True),
        ("Heading 2", 12, INK, True),
        ("Heading 3", 11, INK, True),
        ("List Bullet", 10.5, INK, False),
    ]:
        style = doc.styles[name]
        set_font(style, size, color, bold)
        style.paragraph_format.space_before = Pt(10 if name.startswith("Heading") else 0)
        style.paragraph_format.space_after = Pt(5)
        style.paragraph_format.keep_with_next = name != "List Bullet"
    doc.styles["Heading 1"].paragraph_format.space_after = Pt(12)
    for name, size in [("Contact", 9.5), ("Detail", 9.5), ("Folio", 8.5)]:
        style = doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        set_font(style, size, MUTED)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.keep_with_next = True
    bullet_style = doc.styles["List Bullet"].paragraph_format
    bullet_style.left_indent = Inches(0.14)
    bullet_style.first_line_indent = Inches(-0.14)
    bullet_style.space_after = Pt(4)

    doc.core_properties.title = "Michael Grier Resume"
    doc.core_properties.author = data["name"]
    doc.core_properties.subject = "Software development experience and selected projects"
    doc.core_properties.keywords = "TypeScript, React, Next.js, React Native, software developer"
    doc.add_paragraph(data["name"], "Title")
    doc.add_paragraph(data["title"], "Subtitle")
    contact = doc.add_paragraph(style="Contact")
    contact.add_run(data["location"] + "  |  ")
    link(contact, data["email"], "mailto:" + data["email"])
    contact.add_run("  |  ")
    contact.add_run(data["phone"])
    profiles = doc.add_paragraph(style="Contact")
    profiles.add_run("GitHub: ")
    link(profiles, data["github"].removeprefix("https://"), data["github"])
    profiles.paragraph_format.space_after = Pt(12)
    doc.add_paragraph(data["summary"])
    heading(doc, "Experience")
    for role in data["experience"]:
        doc.add_paragraph(role["title"], "Heading 2")
        doc.add_paragraph(f'{role["organization"]}, {role["location"]}  |  {role["period"]}', "Detail")
        for product in role["products"]:
            paragraph = doc.add_paragraph(style="Heading 3")
            paragraph.add_run(product["name"])
            detail = paragraph.add_run("  |  " + product["description"])
            detail.bold = False
            detail.font.size = Pt(9.5)
            detail.font.color.rgb = RGBColor.from_string(MUTED)
            for bullet in product["bullets"]:
                doc.add_paragraph(bullet, "List Bullet")
        for bullet in role["bullets"]:
            doc.add_paragraph(bullet)

    # A deliberate page boundary keeps the professional and independent work distinct.
    projects_heading = heading(doc, "Selected projects")
    projects_heading.paragraph_format.page_break_before = True
    doc.add_paragraph("Independent development", "Detail")
    for project in data["projects"]:
        paragraph = doc.add_paragraph(style="Heading 2")
        paragraph.paragraph_format.space_before = Pt(12)
        paragraph.paragraph_format.space_after = Pt(2)
        link(paragraph, project["name"], project["href"])
        detail = doc.add_paragraph(project["description"], "Detail")
        detail.paragraph_format.space_after = Pt(7)
        doc.add_paragraph(project["bullet"])
        doc.add_paragraph(project["stack"], "Detail")
    heading(doc, "Technical skills")
    for group in data["skills"]:
        paragraph = doc.add_paragraph()
        paragraph.add_run(group["label"] + ": ").bold = True
        paragraph.add_run(", ".join(group["items"]))
        paragraph.paragraph_format.space_after = Pt(7)
    heading(doc, "Education")
    for education in data["education"]:
        paragraph = doc.add_paragraph()
        paragraph.add_run(education["qualification"]).bold = True
        paragraph.add_run("\n" + education["institution"] + "  |  " + education["detail"])
        paragraph.paragraph_format.space_after = Pt(8)
        paragraph.paragraph_format.keep_together = True

    header = section.header.paragraphs[0]
    header.style = doc.styles["Folio"]
    header.add_run("Michael Grier  /  Software developer")
    lightburst(section.first_page_header.paragraphs[0], 4.0)
    lightburst(header, 2.7)
    for footer in [section.first_page_footer, section.footer]:
        paragraph = footer.paragraphs[0]
        paragraph.style = doc.styles["Folio"]
        paragraph.paragraph_format.keep_with_next = False
        paragraph.add_run("Michael Grier")
        paragraph.paragraph_format.tab_stops.add_tab_stop(Inches(6.95))
        paragraph.add_run("\t")
        field = OxmlElement("w:fldSimple")
        field.set(qn("w:instr"), "PAGE")
        paragraph._p.append(field)
    embed_fonts(doc)
    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)
    print(output)


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output", type=Path, default=ROOT / "public/Michael_Grier_Resume.docx")
    build(parser.parse_args().output)
